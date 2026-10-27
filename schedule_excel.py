# -*- coding: utf-8 -*-
"""openpyxl 기반 Excel 출력."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from schedule_models import FIXED, STL, SUPPORT, UNASSIGNED
from schedule_utils import date_to_str, get_month_dates, is_weekend_or_holiday, manager_display_name, weekday_kr


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WEEKEND_FILL = PatternFill("solid", fgColor="FFF2CC")
UNASSIGNED_FILL = PatternFill("solid", fgColor="F4CCCC")
FIXED_FILL = PatternFill("solid", fgColor="D9EAD3")
STL_FILL = PatternFill("solid", fgColor="D9EAF7")
SUPPORT_FILL = PatternFill("solid", fgColor="EADCF8")
THIN_BORDER = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)


def _style_header(row):
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def _style_table_cell(cell):
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def _assignment_display(assignment, managers):
    if not assignment:
        return "미배정"
    worker_id = assignment.get("worker_id")
    if assignment.get("type") == UNASSIGNED:
        return "미배정"
    if assignment.get("type") == SUPPORT:
        return "외부지원"
    return manager_display_name(worker_id, managers)


def _apply_assignment_fill(cell, assignment):
    if not assignment or assignment.get("type") == UNASSIGNED:
        cell.fill = UNASSIGNED_FILL
    elif assignment.get("type") == FIXED:
        cell.fill = FIXED_FILL
    elif assignment.get("type") == STL:
        cell.fill = STL_FILL
    elif assignment.get("type") == SUPPORT:
        cell.fill = SUPPORT_FILL


def _write_schedule_table(ws, schedule, year, month, store, shifts, managers, holidays):
    headers = ["날짜", "요일"] + shifts
    ws.append(headers)
    _style_header(ws[1])

    row_idx = 2
    for work_date in get_month_dates(year, month):
        date_str = date_to_str(work_date)
        row_values = [date_str, weekday_kr(work_date)]
        for shift in shifts:
            row_values.append(_assignment_display(schedule[date_str][store][shift], managers))
        ws.append(row_values)

        is_special_day = is_weekend_or_holiday(work_date, holidays)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            _style_table_cell(cell)
            if is_special_day:
                cell.fill = WEEKEND_FILL

        for offset, shift in enumerate(shifts, start=3):
            assignment = schedule[date_str][store][shift]
            _apply_assignment_fill(ws.cell(row=row_idx, column=offset), assignment)

        row_idx += 1


def _write_stats(ws, start_col, work_counts, managers):
    headers = ["매니저", "전체", "오픈", "마감", "데일리", "나이트", "주말/공휴일"]
    for idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    row_idx = 2
    for worker_id in managers:
        values = [
            manager_display_name(worker_id, managers),
            work_counts[worker_id]["전체"],
            work_counts[worker_id]["오픈"],
            work_counts[worker_id]["마감"],
            work_counts[worker_id]["데일리"],
            work_counts[worker_id]["나이트"],
            work_counts[worker_id]["주말/공휴일"],
        ]
        for col_idx, value in enumerate(values, start=start_col):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            _style_table_cell(cell)
        row_idx += 1


def _write_issues(ws, start_row, issues, notes):
    ws.cell(row=start_row, column=1, value="검증/예외 사항")
    ws.cell(row=start_row, column=1).font = Font(bold=True)
    row_idx = start_row + 1

    combined = []
    combined.extend(notes or [])
    combined.extend(issues or [])
    if not combined:
        combined = ["검증 결과 특이사항 없음"]

    for item in combined:
        ws.cell(row=row_idx, column=1, value=f"- {item}")
        ws.cell(row=row_idx, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        row_idx += 1


def _set_column_widths(ws):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 34)


def _write_option_summary(wb, schedule_options, selected_option_index):
    """여러 후보 스케줄의 점수표를 별도 시트로 기록한다."""
    ws = wb.create_sheet(title="후보요약")
    headers = ["순위", "후보번호", "선택", "점수", "미배정", "규칙위반", "쏠림", "전략"]
    ws.append(headers)
    _style_header(ws[1])

    for rank, option in enumerate(schedule_options, start=1):
        strategy = option.get("strategy", {})
        strategy_text = (
            f"store_order={strategy.get('store_order')}; "
            f"candidate_offset={strategy.get('candidate_offset')}; "
            f"name={strategy.get('name')}"
        )
        ws.append(
            [
                rank,
                option["index"],
                "선택" if option["index"] == selected_option_index else "",
                option["score"],
                option["unassigned_count"],
                option["violation_count"],
                option["imbalance_score"],
                strategy_text,
            ]
        )
        for cell in ws[ws.max_row]:
            _style_table_cell(cell)
            if option["index"] == selected_option_index:
                cell.fill = FIXED_FILL

    _set_column_widths(ws)


def create_excel(
    schedule,
    work_counts,
    validation_issues,
    notes,
    year,
    month,
    stores,
    managers,
    holidays,
    output_file,
    schedule_options=None,
    selected_option_index=None,
):
    """매장별 3개 시트를 가진 Excel 파일을 생성한다."""
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    for store, shifts in stores.items():
        ws = wb.create_sheet(title=store)
        _write_schedule_table(ws, schedule, year, month, store, shifts, managers, holidays)
        stats_col = len(shifts) + 5
        _write_stats(ws, stats_col, work_counts, managers)
        issue_start_row = len(get_month_dates(year, month)) + 4
        _write_issues(ws, issue_start_row, validation_issues, notes)
        ws.freeze_panes = "A2"
        _set_column_widths(ws)

    if schedule_options:
        _write_option_summary(wb, schedule_options, selected_option_index)

    wb.save(output_file)
